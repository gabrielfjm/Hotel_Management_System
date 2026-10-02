"""Gera as apresentações a partir das evidências.

Uso: uv run --no-project --python 3.12 --with python-pptx==1.0.2 python scripts/gerar_slides.py

Saídas:
    docs/apresentacao-2-slides.pptx   os 2 slides exigidos (processo; resultados, dificuldades e lições),
                                      com notas do apresentador
    docs/apresentacao-apoio-demo.pptx slides de apoio para a demonstração (opcionais)
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
EVID = ROOT / "evidencias"
TELAS = EVID / "telas"
TELAS_VV = EVID / "telas-vvtestlab"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
VINHO = RGBColor(0x8A, 0x10, 0x38)
AZUL = RGBColor(0x1F, 0x5F, 0xA8)
VERDE = RGBColor(0x1E, 0x8E, 0x5A)
LARANJA = RGBColor(0xD9, 0x7A, 0x1E)
CINZA = RGBColor(0x5B, 0x64, 0x72)
CLARO = RGBColor(0xF3, 0xF5, 0xF9)
BORDA = RGBColor(0xD7, 0xDD, 0xE6)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
SUAVE = RGBColor(0xC9, 0xD4, 0xE5)
FONTE = "Segoe UI"
RODAPE = "Gabriel Felipe Jess Meira · Verificação e Validação · github.com/gabrielfjm/Hotel_Management_System"


def pct(v):
    return f"{v:.1f}%".replace(".", ",")


def retangulo(slide, x, y, w, h, cor=None, borda=None, forma=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if cor is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = cor
    if borda is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borda
        s.line.width = Pt(0.75)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    return s


def texto(forma, conteudo, tam=12, cor=NAVY, negrito=False, alinhamento=PP_ALIGN.LEFT, novo=True, espaco=2,
          fonte=FONTE):
    tf = forma.text_frame
    p = tf.add_paragraph() if novo else tf.paragraphs[0]
    p.alignment = alinhamento
    p.space_after = Pt(espaco)
    for i, parte in enumerate(conteudo.split("**")):
        if not parte:
            continue
        r = p.add_run()
        r.text = parte
        r.font.size = Pt(tam)
        r.font.name = fonte
        r.font.color.rgb = cor
        r.font.bold = negrito or i % 2 == 1
    return p


def imagem(slide, caminho, x, y, w=None, h=None, borda=True):
    if not caminho.exists():
        raise FileNotFoundError(caminho)
    img = slide.shapes.add_picture(str(caminho), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if borda:
        img.line.color.rgb = BORDA
        img.line.width = Pt(1)
    return img


def cabecalho(slide, numero, titulo, subtitulo):
    retangulo(slide, 0, 0, 13.333, 1.05, NAVY)
    retangulo(slide, 0, 1.05, 13.333, 0.06, VINHO)
    selo = retangulo(slide, 0.45, 0.22, 0.62, 0.62, VINHO, forma=MSO_SHAPE.OVAL)
    selo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    selo.text_frame.margin_left = selo.text_frame.margin_right = 0
    texto(selo, str(numero), 20, BRANCO, True, PP_ALIGN.CENTER, novo=False)
    texto(retangulo(slide, 1.25, 0.1, 11.6, 0.55), titulo, 25, BRANCO, True, novo=False)
    texto(retangulo(slide, 1.25, 0.6, 11.6, 0.4), subtitulo, 12, SUAVE, novo=False)
    texto(retangulo(slide, 0.45, 7.12, 12.4, 0.3), RODAPE, 9, CINZA, novo=False)


def legenda(slide, x, y, w, conteudo):
    texto(retangulo(slide, x, y, w, 0.3), conteudo, 9.5, CINZA, alinhamento=PP_ALIGN.CENTER, novo=False)


def notas(slide, conteudo):
    slide.notes_slide.notes_text_frame.text = conteudo


# ---------------------------------------------------------------- slides principais

EQUIV = 7  # sobreviventes equivalentes (S8 a S14, relatório 7.2)


def escore(mortos, total):
    """Escore de mutação da disciplina: mortos ÷ (gerados − equivalentes) × 100."""
    return 100 * mortos / (total - EQUIV)


def escore_evolucao(linha):
    """Escore de uma linha de evolucao.json ("127/141"), ou None nas etapas sem mutação."""
    if linha["escore_pct"] is None:
        return None
    mortos, total = (int(n) for n in linha["mutacao"].split("/"))
    return escore(mortos, total)


def slide_processo(prs, evol, mut_ini, mut_fim):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cabecalho(s, 1, "Processo de desenvolvimento do teste",
              "Hotel Management System (Flask + SQLite) · REQ-01 Reservar quartos · REQ-02 Data de entrada · REQ-03 Data de saída")
    e = {linha["etapa"][0]: linha for linha in evol}
    etapas = [
        ("Funcional", "pytest · só a especificação", AZUL,
         ["**22 classes** (11 válidas, 11 inválidas)", "1 caso por classe inválida, limites embutidos",
          f'**{e["1"]["casos"]} casos** no código original', f'**{e["1"]["xfail"]} falham** → 9 defeitos']),
        ("Estrutural", "coverage.py", AZUL,
         ["Meta: **100%** de comandos e desvios viáveis", '**5 casos reaproveitados**; 3 ampliados nas lacunas',
          "**+1 defeito** só visível no código", f'**{pct(e["2"]["pct_desvios"])}** dos desvios (1 inviável)']),
        ("Correção", "fork no GitHub · 6 commits", VINHO,
         ["**10 defeitos** corrigidos (+9 dos secundários)", "Cada commit confirmado pelo teste que revelou o defeito",
          f'**{e["3"]["casos"]}/{e["3"]["casos"]}** casos passam, sem skip', "**100%** de comandos e desvios"]),
        ("Mutação", f'Cosmic Ray · {mut_fim["total"]} mutantes', VERDE,
         [f'Escore = mortos ÷ (gerados − {EQUIV} equivalentes)',
          f'Inicial: **{pct(escore(mut_ini["mortos"], mut_ini["total"]))}** ({mut_ini["mortos"]}/{mut_ini["total"] - EQUIV})',
          f'{mut_ini["sobreviventes"]} vivos: **{mut_ini["sobreviventes"] - EQUIV} matáveis**; **4 casos ampliados**',
          f'Final: **{pct(escore(mut_fim["mortos"], mut_fim["total"]))}** ({mut_fim["mortos"]}/{mut_fim["total"] - EQUIV})']),
    ]
    x = 0.45
    for i, (nome, ferramenta, cor, itens) in enumerate(etapas):
        seta = retangulo(s, x, 1.4, 2.02, 0.74, cor, forma=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        seta.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        seta.text_frame.margin_left = Inches(0.3 if i else 0.15)
        texto(seta, nome, 13, BRANCO, True, novo=False, espaco=0)
        texto(seta, ferramenta, 8.5, BRANCO, espaco=0)
        card = retangulo(s, x + 0.03, 2.25, 1.86, 2.6, CLARO, BORDA)
        for j, item in enumerate(itens):
            texto(card, item, 10.5, NAVY, novo=j > 0, espaco=6)
        x += 1.96
    cadeia = retangulo(s, 0.45, 5.02, 7.77, 0.95, BRANCO, BORDA)
    texto(cadeia, "Rastreabilidade de ponta a ponta", 11.5, VINHO, True, novo=False)
    texto(cadeia, "**REQ** → **CE** (@pytest.mark.ce) → **CT** (nome do teste) → **execução** (JUnit) → "
                  "**DEF** (@pytest.mark.defeito) → **commit** de correção", 10.5, NAVY)
    texto(retangulo(s, 0.45, 6.05, 7.77, 1.0),
          "Mesmo teste nos dois códigos: **xfail estrito no original** (o defeito precisa aparecer) e **verde no "
          "corrigido**. etapas.ps1 / etapas.sh reproduzem tudo e gravam as evidências de cada etapa.",
          10.5, CINZA, novo=False)
    imagem(s, TELAS / "original-03-quartos.png", 8.5, 1.4, w=4.4)
    legenda(s, 8.5, 3.34, 4.4, "SUT em execução: lista de quartos após o login")
    img = imagem(s, TELAS_VV / "04-rastreabilidade.png", 8.5, 3.78, w=4.4)
    legenda(s, 8.5, 3.78 + Emu(img.height).inches + 0.02, 4.4,
            "V&V TestLab: requisito → classe → caso → execução → defeito")
    notas(s, (
        "Objetivo (30 s): testar um sistema web de terceiros, o Hotel Management System em Flask, em três requisitos "
        "da reserva de quartos: reservar (quartos, hóspedes, disponibilidade, custo), data de entrada e data de saída; os demais RFs foram documentados.\n\n"
        f"Funcional (40 s): só com o README e o formulário de reserva, particionei 10 condições em 22 classes: "
        f"um caso para cada classe inválida, com os limites embutidos (entrada hoje, uma noite, capacidade, estadias adjacentes). "
        f"Foram {e['1']['casos']} casos; no código original {e['1']['xfail']} falharam e revelaram 9 defeitos.\n\n"
        "Estrutural (40 s): medi com coverage.py --branch e defini a meta de 100% dos desvios viáveis. Os ramos que "
        "faltavam foram cobertos sem criar casos: reaproveitei 5 funcionais para percorrer os grafos e ampliei 3 deles. Lendo o laço triplo, achei um defeito que a especificação não sugeria: "
        "um produto cartesiano entre reservas e vínculos, mascarado pela tautologia.\n\n"
        "Correção (30 s): antes de mutar, corrigi os defeitos no fork, um commit por grupo. O mesmo teste é xfail "
        "estrito no original e passa no corrigido: isso comprova cada correção.\n\n"
        f"Mutação (40 s): Cosmic Ray sobre todos os {mut_fim['total']} mutantes das funções dos três requisitos. Dos {mut_ini['sobreviventes']} vivos, {mut_ini['sobreviventes'] - EQUIV} eram matáveis "
        f"e foram mortos ampliando 4 dos mesmos 15 casos; os {EQUIV} restantes são equivalentes e estão justificados no relatório. "
        f"Pela fórmula da disciplina, que tira os equivalentes do denominador, o escore foi de "
        f"{pct(escore(mut_ini['mortos'], mut_ini['total']))} para {pct(escore(mut_fim['mortos'], mut_fim['total']))} "
        f"(sem descontar: {pct(mut_ini['escore_pct'])} e {pct(mut_fim['escore_pct'])}).\n\n"
        "Ferramenta (20 s): organizei tudo no V&V TestLab, que importa JUnit, cobertura e mutação pela ponte local e "
        "mantém a rastreabilidade até o defeito."))


def slide_resultados(prs, evol, mut_fim):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    final = evol[-1]
    cabecalho(s, 2, "Resultados, dificuldades e lições aprendidas",
              f'os mesmos {final["casos"]} casos nas 3 etapas · 10 defeitos, todos corrigidos · 100% dos desvios · '
              f'{pct(escore(mut_fim["mortos"], mut_fim["total"]))} de escore de mutação')
    x = 0.45
    for valor, rotulo, cor in [(str(final["casos"]), "casos de teste", AZUL), ("10", "defeitos corrigidos", VINHO),
                               ("100%", "dos desvios cobertos", VERDE),
                               (pct(escore(mut_fim["mortos"], mut_fim["total"])), f"escore de mutação ({mut_fim['mortos']} ÷ {mut_fim['total'] - EQUIV})", LARANJA)]:
        k = retangulo(s, x, 1.3, 1.9, 0.95, CLARO, BORDA)
        k.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        texto(k, valor, 24, cor, True, PP_ALIGN.CENTER, novo=False, espaco=0)
        texto(k, rotulo, 9.5, CINZA, alinhamento=PP_ALIGN.CENTER, espaco=0)
        x += 2.0
    dados = CategoryChartData()
    dados.categories = [f'{l["etapa"].split(". ")[1]} ({l["casos"]})' for l in evol]
    dados.add_series("Comandos", [l["pct_comandos"] / 100 for l in evol])
    dados.add_series("Desvios", [l["pct_desvios"] / 100 for l in evol])
    dados.add_series("Escore de mutação", [escore_evolucao(l) / 100 if l["escore_pct"] is not None else None for l in evol])
    grafico = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.45), Inches(2.38), Inches(7.9),
                                 Inches(2.62), dados).chart
    grafico.has_legend = True
    grafico.legend.position = XL_LEGEND_POSITION.TOP
    grafico.legend.include_in_layout = False
    grafico.legend.font.size = Pt(9)
    grafico.legend.font.name = FONTE
    for serie, cor in zip(grafico.series, (AZUL, VINHO, VERDE)):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = cor
        rotulos = serie.data_labels
        rotulos.show_value = True
        rotulos.number_format = "0%"
        rotulos.number_format_is_linked = False
        rotulos.font.size = Pt(8)
        rotulos.font.name = FONTE
        rotulos.position = XL_LABEL_POSITION.OUTSIDE_END
    eixo = grafico.value_axis
    eixo.maximum_scale = 1.18
    eixo.minimum_scale = 0
    eixo.visible = False
    eixo.has_major_gridlines = False
    grafico.category_axis.tick_labels.font.size = Pt(9)
    grafico.category_axis.tick_labels.font.name = FONTE
    grafico.plots[0].gap_width = 70
    texto(retangulo(s, 8.6, 1.28, 4.3, 0.32), "DEF-01 · reservar o 101 fora do período ocupado", 10.5, VINHO, True,
          novo=False)
    y = 1.62
    for arquivo, rotulo, cor in [("original-07-reserva-resultado.png", "ORIGINAL · recusada", VINHO),
                                 ("corrigida-07-reserva-resultado.png", "CORRIGIDO · aceita", VERDE)]:
        img = imagem(s, TELAS / arquivo, 8.6, y, w=4.3)
        img.crop_bottom = 0.5
        img.height = int(img.height * 0.5)
        etiqueta = retangulo(s, 11.25, y + Emu(img.height).inches - 0.32, 1.6, 0.26, cor)
        etiqueta.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        texto(etiqueta, rotulo, 8.5, BRANCO, True, PP_ALIGN.CENTER, novo=False, espaco=0)
        y += Emu(img.height).inches + 0.08
    defeitos = retangulo(s, 8.6, y + 0.04, 4.3, 5.0 - (y + 0.04), CLARO, BORDA)
    texto(defeitos, "Defeitos mais graves", 11, VINHO, True, novo=False, espaco=2)
    for item in ["**Tautologia** no conflito de datas bloqueia o quarto",
                 "**Entrada hoje** recusada (data comparada com a hora)",
                 "**Quarto inexistente ou repetido** aceito na reserva",
                 "Defeito **mascarado** (produto cartesiano)"]:
        texto(defeitos, "• " + item, 9.5, NAVY, espaco=1)
    dif = retangulo(s, 0.45, 5.15, 4.0, 1.9, BRANCO, BORDA)
    texto(dif, "Dificuldades", 12, VINHO, True, novo=False, espaco=3)
    for item in ["Dependências de 2019: Python 3.9 e Python 3.12 separados",
                 "Cosmic Ray no Windows (caminho do interpretador)",
                 "Oráculos não escritos → convenções declaradas"]:
        texto(dif, "• " + item, 10, NAVY, espaco=2)
    lic = retangulo(s, 4.6, 5.15, 8.3, 1.9, BRANCO, BORDA)
    texto(lic, "Lições aprendidas", 12, VINHO, True, novo=False, espaco=3)
    falhas = evol[1]["xfail"]
    for item in ["**Limites valem mais que valores típicos**: 15 casos funcionais acharam 9 dos 10 defeitos",
                 f"**Cobertura ≠ correção**: 100% dos comandos no original com {falhas} testes falhando",
                 f"**Mutação testa os testes**: com 100% dos desvios ainda havia 7 mutantes não equivalentes vivos",
                 "**Corrigir antes de mutar**: mutar código com defeitos invalida o escore"]:
        texto(lic, "• " + item, 10, NAVY, espaco=2)
    notas(s, (
        f"Números (30 s): os mesmos {final['casos']} casos atravessam as três etapas (ampliados quando faltou um cenário), com 100% de comandos e desvios nas funções do recorte "
        f"e {pct(escore(mut_fim['mortos'], mut_fim['total']))} de escore de mutação: {mut_fim['mortos']} mortos ÷ "
        f"({mut_fim['total']} gerados − {EQUIV} equivalentes). Os {EQUIV} mutantes vivos são equivalentes, por isso saem do "
        f"denominador (sem descontar, o escore seria {pct(mut_fim['escore_pct'])}). No gráfico, a cobertura sobe na etapa "
        "estrutural e o escore sobe na mutação.\n\n"
        "Defeito principal (40 s): a condição de conflito de datas é uma tautologia, sempre verdadeira. Depois da "
        "primeira reserva, o quarto nunca mais pode ser reservado. As imagens mostram a mesma reserva recusada no "
        "original e aceita no corrigido. Outro exemplo: a entrada no dia de hoje era recusada, porque a data era comparada com a hora atual.\n\n"
        "Defeito mascarado (20 s): o laço comparava cada quarto com reservas de outros quartos; isso só apareceu depois "
        "de corrigir a tautologia.\n\n"
        "Dificuldades (20 s): o projeto é de 2019 e não roda no Python atual; usei dois ambientes. O Cosmic Ray no "
        "Windows exigiu caminho absoluto do interpretador. Onde o README não definia o comportamento, declarei convenções.\n\n"
        f"Lições (40 s): limites acham defeitos que valores típicos não acham; cobertura mede o que executou, não o que "
        f"foi verificado (o original tinha 100% dos comandos com {falhas} testes falhando); a mutação avaliou a própria "
        "suíte e mostrou fragilidades mesmo com 100% dos desvios; e é preciso corrigir antes de mutar."))


# ------------------------------------------------------------------ deck de apoio

def slide_titulo(prs, titulo, linhas):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    retangulo(s, 0, 0, 13.333, 7.5, NAVY)
    retangulo(s, 0.8, 3.55, 1.4, 0.08, VINHO)
    texto(retangulo(s, 0.8, 1.9, 11.5, 1.6), titulo, 34, BRANCO, True, novo=False)
    st = retangulo(s, 0.8, 3.8, 11.5, 1.8)
    for i, linha in enumerate(linhas):
        texto(st, linha, 15, SUAVE, novo=i > 0, espaco=4)


def slide_grade(prs, numero, titulo, subtitulo, imagens, nota):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cabecalho(s, numero, titulo, subtitulo)
    for (caminho, rotulo), (x, y) in zip(imagens, [(0.45, 1.35), (6.75, 1.35), (0.45, 4.2), (6.75, 4.2)]):
        img = imagem(s, caminho, x, y, h=2.4)
        if Emu(img.width).inches > 6.1:
            proporcao = Inches(6.1) / img.width
            img.width, img.height = Inches(6.1), int(img.height * proporcao)
        legenda(s, x, y + Emu(img.height).inches + 0.02, Emu(img.width).inches, rotulo)
    notas(s, nota)


def slide_codigo(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cabecalho(s, "D", "DEF-01: o predicado de conflito é uma tautologia",
              "hotel/views.py · reserve() e show_rooms() · commit 2d1b2cd no fork")
    antes = retangulo(s, 0.45, 1.4, 6.1, 3.15, RGBColor(0xFD, 0xEE, 0xEE), RGBColor(0xE8, 0xB4, 0xB4))
    texto(antes, "Original (linhas 208–210)", 12, VINHO, True, novo=False)
    for linha in ["if each_booking.room_id == int(each) and (",
                  "    (c1 <= d1 and d2 <= c2) or (d1 <= c1 and d2 <= c2) or",
                  "    (c1 <= d1 and c2 <= d2) or (d1 <= c1 and c2 <= d2)):",
                  '    flash("...not available...")']:
        texto(antes, linha, 10.5, NAVY, fonte="Consolas")
    texto(antes, "Se d1 ≤ c1, vale a 2ª ou a 4ª alternativa (d2 ≤ c2 ou c2 ≤ d2); se c1 ≤ d1, vale a 1ª ou a 3ª. "
                 "Para quaisquer datas, a condição é verdadeira.", 10.5, CINZA)
    depois = retangulo(s, 6.8, 1.4, 6.1, 3.15, RGBColor(0xEA, 0xF7, 0xEF), RGBColor(0xA9, 0xD8, 0xBC))
    texto(depois, "Corrigido", 12, VERDE, True, novo=False)
    for linha in ["def _periodos_conflitam(a1, a2, b1, b2):",
                  "    return a1 < b2 and b1 < a2   # [entrada, saída)",
                  " ",
                  "vinculos = db.session.query(...).join(",
                  "    Reservations, Booked.brid == Reservations.rid)"]:
        texto(depois, linha, 10.5, NAVY, fonte="Consolas")
    texto(depois, "Interseção de intervalos semiabertos + junção de cada vínculo com a própria reserva (DEF-15).",
          10.5, CINZA)
    casos = retangulo(s, 0.45, 4.8, 12.45, 2.2, CLARO, BORDA)
    texto(casos, "Como os testes revelaram", 12, VINHO, True, novo=False)
    for item in ["Funcional (valor limite): CT-003 (entrar em 12/03/2030, dia em que a outra estadia sai) foi recusado no original",
                 "Estrutural: a ampliação do CT-012 revelou DEF-15 (produto cartesiano), escondido pela tautologia até ela ser corrigida",
                 "Mutação: a ampliação do CT-003 (sair em 10/03, dia em que a outra entra; estadia antes) matou 3 mutantes do predicado",
                 "Mutação: a ampliação do CT-012 matou o mutante Booked.brid >= Reservations.rid na junção corrigida"]:
        texto(casos, "• " + item, 10.5, NAVY, espaco=3)
    notas(s, "Use se perguntarem como um defeito tão grave passou despercebido: com uma única reserva no banco o "
             "sistema parece funcionar; os valores limite de estadias adjacentes e períodos distantes expõem o problema.")


def slide_roteiro(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cabecalho(s, "R", "Roteiro da demonstração ao vivo", "~5 minutos · dois terminais e o navegador")
    passos = [
        ("Sistema", "output/hotel-management-testado → .\\executar.ps1 → http://127.0.0.1:5000 (ana@example.test / senha123)"),
        ("Suíte", ".\\etapas.ps1 -SemMutacao → etapas 1 e 2 no original (defeitos em xfail) e suíte final toda verde"),
        ("V&V TestLab", "raiz → npm run dev → Dados e exportação → importar output/hotel-vvtestlab-projeto-inicial.json"),
        ("Ponte no original", ".\\iniciar-integracao.ps1 -Versao original → Integração Python → Executar e sincronizar: falhas viram defeitos"),
        ("Ponte no corrigido", "Ctrl+C e .\\iniciar-integracao.ps1 → Executar e sincronizar: todos aprovados, 100% dos desvios"),
        ("Mutação", "Executar mutação (~4 min) → 134/141 mortos; página Teste de mutação mostra o cálculo"),
    ]
    y = 1.4
    for n, (titulo, detalhe) in enumerate(passos, 1):
        selo = retangulo(s, 0.55, y, 0.55, 0.55, AZUL, forma=MSO_SHAPE.OVAL)
        selo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        selo.text_frame.margin_left = selo.text_frame.margin_right = 0
        texto(selo, str(n), 16, BRANCO, True, PP_ALIGN.CENTER, novo=False)
        caixa = retangulo(s, 1.3, y - 0.05, 11.5, 0.8)
        texto(caixa, titulo, 13, VINHO, True, novo=False, espaco=0)
        texto(caixa, detalhe, 11, NAVY, espaco=0)
        y += 0.93
    notas(s, "Se faltar tempo, pule a mutação ao vivo e importe o estudo completo (output/hotel-vvtestlab-backup.json), "
             "que já traz execuções, defeitos e métricas.")


def main():
    evol = json.loads((EVID / "evolucao.json").read_text(encoding="utf-8"))
    mut_ini = json.loads((EVID / "mutacao-inicial" / "resumo.json").read_text(encoding="utf-8"))
    mut_fim = json.loads((EVID / "mutacao-final" / "resumo.json").read_text(encoding="utf-8"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_processo(prs, evol, mut_ini, mut_fim)
    slide_resultados(prs, evol, mut_fim)
    prs.save(ROOT / "docs" / "apresentacao-2-slides.pptx")

    apoio = Presentation()
    apoio.slide_width, apoio.slide_height = Inches(13.333), Inches(7.5)
    slide_titulo(apoio, "Teste de software de terceiros: Hotel Management System",
                 ["Slides de apoio à demonstração", "Gabriel Felipe Jess Meira · Verificação e Validação",
                  "github.com/gabrielfjm/Hotel_Management_System"])
    slide_grade(apoio, "S", "O sistema em execução", "Flask 0.12 + SQLite · capturas do código original e do corrigido", [
        (TELAS / "original-02-login.png", "Login (original)"),
        (TELAS / "original-04-consulta-formulario.png", "Consulta de disponibilidade: 20/10/2026 a 22/10/2026, 4 hóspedes"),
        (TELAS / "original-05-consulta-resultado.png", "Original: o 101 livre some (DEF-01) e o 102, de capacidade 3, aparece (DEF-14)"),
        (TELAS / "corrigida-05-consulta-resultado.png", "Corrigido: só o quarto de capacidade 4 aparece"),
    ], "Mostre a mesma consulta nos dois códigos: a diferença resume dois defeitos.")
    slide_codigo(apoio)
    slide_grade(apoio, "V", "V&V TestLab: gestão do projeto de testes",
                "estudo do hotel importado do backup gerado a partir das evidências", [
        (TELAS_VV / "01-visao-geral.png", "Visão geral das três etapas"),
        (TELAS_VV / "09-teste-funcional-detalhe.png", "Teste funcional: entrada, esperado, obtido e assertivas"),
        (TELAS_VV / "10-teste-estrutural-grafo.png", "Teste estrutural: grafo percorrido pelo caso"),
        (TELAS_VV / "11-teste-mutacao.png", "Teste de mutação: cálculo do escore e mutantes"),
    ], "A ferramenta importa JUnit, cobertura e resultado do Cosmic Ray pela ponte local e vincula tudo pelo ID CT-xxx.")
    slide_roteiro(apoio)
    apoio.save(ROOT / "docs" / "apresentacao-apoio-demo.pptx")
    print("docs/apresentacao-2-slides.pptx e docs/apresentacao-apoio-demo.pptx")


if __name__ == "__main__":
    main()
